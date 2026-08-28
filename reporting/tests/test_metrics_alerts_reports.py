import tempfile
from datetime import date, timedelta
from decimal import Decimal
from unittest.mock import patch

from django.test import TestCase, override_settings
from openpyxl import load_workbook

from reporting.models import (
    AdAccount,
    Anomaly,
    InsightDaily,
    InsightLevel,
    MetaConnection,
    RawApiPayload,
    ReportRun,
    ReportScope,
    ReportStatus,
    Severity,
    SyncRun,
    SyncStatus,
)
from reporting.services.alerts import evaluate_alerts
from reporting.services.metrics import aggregate_rows
from reporting.services.reports import generate_portfolio_report, generate_report


class MetricsAlertsAndReportsTests(TestCase):
    def setUp(self):
        self.connection = MetaConnection.objects.create(name="Meta", ad_account_external_id="456", api_version="v25.0")
        self.account = AdAccount.objects.create(
            connection=self.connection,
            external_id="456",
            name="ULTEx Test",
            currency="MAD",
            timezone_name="Africa/Casablanca",
        )

    def _insight(self, day, *, results=Decimal("10"), spend=Decimal("100"), level=InsightLevel.ACCOUNT, object_id="456", verified=True, raw=None, sync=None):
        return InsightDaily.objects.create(
            account=self.account,
            sync_run=sync,
            raw_payload=raw,
            level=level,
            object_external_id=object_id,
            object_name="ULTEx Test" if level == InsightLevel.ACCOUNT else "Campagne",
            date=day,
            currency="MAD",
            spend=spend,
            impressions=1000,
            reach=800,
            clicks=40,
            link_clicks=30,
            frequency=Decimal("1.25"),
            ctr=Decimal("4"),
            results=results,
            cost_per_result=(spend / results) if results else None,
            result_action_type="onsite_conversion.lead_grouped" if results is not None else "",
            result_verified=verified,
        )

    def test_missing_results_propagate_instead_of_becoming_zero(self):
        day = date(2026, 8, 10)
        self._insight(day, results=None)
        aggregate = aggregate_rows(InsightDaily.objects.all())
        self.assertIsNone(aggregate["results"])
        self.assertIsNone(aggregate["cost_per_result"])

    def test_missing_metric_creates_information_alert_not_zero_result_alert(self):
        target = date(2026, 8, 10)
        for offset in (3, 2, 1):
            self._insight(target - timedelta(days=offset))
        self._insight(target, results=None, spend=Decimal("200"))
        evaluate_alerts(self.account, target)
        rules = set(Anomaly.objects.filter(date=target).values_list("rule_id", flat=True))
        self.assertIn("METRIC_UNAVAILABLE", rules)
        self.assertNotIn("ZERO_RESULTS_HIGH_SPEND", rules)

    def test_zero_results_threshold_requires_history_and_raises_high_alert(self):
        target = date(2026, 8, 10)
        for offset in (3, 2, 1):
            self._insight(target - timedelta(days=offset))
        self._insight(target, results=Decimal("0"), spend=Decimal("20"))
        evaluate_alerts(self.account, target)
        anomaly = Anomaly.objects.get(date=target, rule_id="ZERO_RESULTS_HIGH_SPEND")
        self.assertEqual(anomaly.severity, Severity.HIGH)

    def _report_fixture(self, day):
        sync = SyncRun.objects.create(
            connection=self.connection,
            account=self.account,
            requested_start=day,
            requested_end=day,
            levels=["account"],
            status=SyncStatus.SUCCESS,
            records_count=1,
            raw_pages_count=1,
        )
        raw = RawApiPayload.objects.create(
            sync_run=sync,
            endpoint="act_456/insights",
            request_params={"level": "account"},
            payload={"data": [{"date_start": day.isoformat()}]},
            checksum="a" * 64,
        )
        self._insight(day, spend=Decimal("123.45"), results=Decimal("9"), raw=raw, sync=sync)

    def test_pdf_and_excel_share_one_immutable_snapshot(self):
        day = date(2026, 8, 12)
        self._report_fixture(day)
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            version = generate_report(self.account, day, day)
            self.assertEqual(version.report_run.status, ReportStatus.READY)
            with version.pdf_file.open("rb") as stream:
                self.assertTrue(stream.read().startswith(b"%PDF"))
            with version.excel_file.open("rb") as stream:
                workbook = load_workbook(stream, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["Résumé", "Tendance quotidienne", "Campagnes", "Ensembles", "Publicités", "Alertes", "Synchronisations"],
            )
            self.assertEqual(Decimal(str(workbook["Résumé"]["B9"].value)), Decimal(version.snapshot["kpis"]["spend"]))
            self.assertEqual(workbook["Résumé"]["D1"].value, "ULTEX · META REPORTS")
            self.assertEqual(len(workbook["Résumé"]._images), 1)
            self.assertEqual(workbook["Tendance quotidienne"]["I2"].value, version.snapshot["trend"][0]["id"])

    def test_portfolio_report_merges_accounts_and_keeps_source_traceability(self):
        day = date(2026, 8, 14)
        self._report_fixture(day)
        self._insight(day, level=InsightLevel.CAMPAIGN, object_id="campaign-first", spend=Decimal("123.45"), results=Decimal("9"))
        second_connection = MetaConnection.objects.create(name="Meta 2", ad_account_external_id="789", api_version="v25.0")
        second_account = AdAccount.objects.create(
            connection=second_connection,
            external_id="789",
            name="ULTEx Second",
            currency="MAD",
            timezone_name="Africa/Casablanca",
        )
        for level, object_id in ((InsightLevel.ACCOUNT, "789"), (InsightLevel.CAMPAIGN, "campaign-second")):
            InsightDaily.objects.create(
                account=second_account,
                level=level,
                object_external_id=object_id,
                object_name="ULTEx Second" if level == InsightLevel.ACCOUNT else "Campagne Second",
                date=day,
                currency="MAD",
                spend=Decimal("50"),
                impressions=500,
                reach=400,
                clicks=20,
                results=Decimal("2"),
                cost_per_result=Decimal("25"),
                result_action_type="onsite_conversion.lead_grouped",
                result_verified=True,
            )

        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            version = generate_portfolio_report([self.account, second_account], day, day)
            self.assertEqual(version.report_run.scope, ReportScope.PORTFOLIO)
            self.assertEqual(version.report_run.included_accounts.count(), 2)
            self.assertTrue(version.snapshot["account"]["combined"])
            self.assertEqual(version.snapshot["account"]["count"], 2)
            self.assertEqual(Decimal(version.snapshot["kpis"]["spend"]), Decimal("173.45"))
            self.assertEqual(Decimal(version.snapshot["kpis"]["results"]), Decimal("11.00"))
            with version.excel_file.open("rb") as stream:
                workbook = load_workbook(stream, data_only=False)
            self.assertEqual(workbook["Campagnes"]["A1"].value, "Compte Meta")
            self.assertEqual({workbook["Campagnes"]["A2"].value, workbook["Campagnes"]["A3"].value}, {"ULTEx Test", "ULTEx Second"})
            with version.pdf_file.open("rb") as stream:
                self.assertTrue(stream.read().startswith(b"%PDF"))

    def test_report_failure_is_persisted_and_raises_critical_alert(self):
        day = date(2026, 8, 13)
        self._report_fixture(day)
        with patch("reporting.services.reports.build_excel", side_effect=RuntimeError("disk failure")):
            with self.assertRaises(RuntimeError):
                generate_report(self.account, day, day)
        run = ReportRun.objects.get(account=self.account, date_start=day, date_end=day)
        self.assertEqual(run.status, ReportStatus.FAILED)
        anomaly = Anomaly.objects.get(account=self.account, date=day, rule_id="REPORT_GENERATION_FAILED")
        self.assertEqual(anomaly.severity, Severity.CRITICAL)
