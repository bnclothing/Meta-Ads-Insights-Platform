from __future__ import annotations

import base64
from datetime import datetime
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from django.conf import settings
from django.template.loader import render_to_string
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.shapes import Drawing, String
from reportlab.platypus import Image as RLImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BRAND = "0159A3"
NAVY = "004472"
PALE = "EAF3FA"
GOLD = "FFC90D"
LINE = "D8E3EC"
LOGO_PATH = Path(settings.BASE_DIR) / "static" / "reporting" / "brand" / "ultex-logo.png"


def _number(value):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return value


def _localized_timestamp(value, timezone_name):
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed.astimezone(ZoneInfo(timezone_name)).strftime("%d/%m/%Y %H:%M")
    except (TypeError, ValueError, ZoneInfoNotFoundError):
        return str(value)


def _logo_data_uri():
    if not LOGO_PATH.exists():
        return ""
    encoded = base64.b64encode(LOGO_PATH.read_bytes()).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def _style_sheet(ws, widths=None, *, landscape=False):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=GOLD))
    ws.row_dimensions[1].height = 24
    thin = Side(style="thin", color=LINE)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if widths:
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = BRAND
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.oddFooter.left.text = "ULTEx - Rapport Meta interne"
    ws.oddFooter.right.text = "Page &[Page] / &[Pages]"


def _add_excel_branding(ws):
    if not LOGO_PATH.exists():
        return
    ws.merge_cells("D1:F1")
    banner = ws["D1"]
    banner.value = "ULTEX · META REPORTS"
    banner.fill = PatternFill("solid", fgColor=GOLD)
    banner.font = Font(color=NAVY, bold=True, size=11)
    banner.alignment = Alignment(horizontal="center", vertical="center")
    banner.border = Border(bottom=Side(style="medium", color=NAVY))
    for column in ("D", "E", "F"):
        ws.column_dimensions[column].width = 13
    logo = XLImage(str(LOGO_PATH))
    logo.width = 245
    logo.height = 81
    ws.add_image(logo, "D2")


def _format_columns(ws, formats):
    for column, number_format in formats.items():
        for cell in ws[column][1:]:
            cell.number_format = number_format


def build_excel(snapshot: dict) -> bytes:
    combined = bool(snapshot["account"].get("combined"))
    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet("Résumé")
    summary.append(["Rubrique", "Valeur"])
    summary_rows = [
        ("Compte", snapshot["account"]["name"]),
        ("Période", f"{snapshot['period']['start']} au {snapshot['period']['end']}"),
        ("Généré le", _localized_timestamp(snapshot["generated_at"], snapshot["account"]["timezone"])),
        ("Devise", snapshot["account"]["currency"]),
        ("Fuseau horaire", snapshot["account"]["timezone"]),
        ("État des données", snapshot["synchronization_status"]),
        ("Résumé exécutif", snapshot["executive_summary"]),
        ("Dépenses", _number(snapshot["kpis"]["spend"])),
        ("Meta résultats", _number(snapshot["kpis"]["results"])),
        ("Coût par résultat", _number(snapshot["kpis"]["cost_per_result"])),
        ("Impressions", snapshot["kpis"]["impressions"]),
        ("Clics", snapshot["kpis"]["clicks"]),
        ("CTR", _number(snapshot["kpis"]["ctr"])),
        ("Empreinte du jeu de données", ", ".join(str(value) for value in snapshot["sources"]["raw_payload_ids"])),
    ]
    for label, level, rank in (
        ("Campagne la plus efficace", "campaigns", "strongest"),
        ("Campagne à examiner", "campaigns", "weakest"),
        ("Publicité la plus efficace", "ads", "strongest"),
        ("Publicité à examiner", "ads", "weakest"),
    ):
        item = snapshot["rankings"][level][rank]
        summary_rows.append((label, f"{item['name']} - {item['cost_per_result']} {snapshot['account']['currency']} / résultat" if item else "Non calculable"))
    for row in summary_rows:
        summary.append(row)
    if combined:
        summary.append(("Comptes inclus", snapshot["account"].get("count", 0)))
        summary.append(
            (
                "Liste des comptes",
                ", ".join(f"{item['name']} (act_{item['external_id']})" for item in snapshot["account"].get("accounts", [])),
            )
        )
    _style_sheet(summary, [28, 90])
    _add_excel_branding(summary)
    summary.row_dimensions[8].height = 62
    _format_columns(summary, {"B": "#,##0.00"})
    for row_number in (2, 3, 4, 5, 6, 7, 8, 15, 16, 17, 18, 19):
        summary[f"B{row_number}"].number_format = "General"
    for row_number in (12, 13):
        summary[f"B{row_number}"].number_format = "#,##0"

    trend = wb.create_sheet("Tendance quotidienne")
    trend.append(["Date", "Dépenses", "Résultats", "Coût / résultat", "Impressions", "Portée", "Clics", "CTR (%)", "Insight ID", "Payload ID"])
    for row in snapshot["trend"]:
        insight_trace = ",".join(str(value) for value in row.get("insight_ids", [])) if combined else row.get("id")
        payload_trace = ",".join(str(value) for value in row.get("raw_payload_ids", [])) if combined else row.get("raw_payload_id")
        trend.append([row["date"], _number(row["spend"]), _number(row["results"]), _number(row["cost_per_result"]), row["impressions"], row["reach"], row["clicks"], _number(row["ctr"]), insight_trace, payload_trace])
    _style_sheet(trend, [14, 14, 12, 18, 14, 14, 12, 12, 12, 12], landscape=True)
    _format_columns(trend, {"B": "#,##0.00", "C": "#,##0.00", "D": "#,##0.00", "E": "#,##0", "F": "#,##0", "G": "#,##0", "H": "0.00"})

    for sheet_name, key in (("Campagnes", "campaigns"), ("Ensembles", "adsets"), ("Publicités", "ads")):
        ws = wb.create_sheet(sheet_name)
        headers = ["Nom", "ID Meta", "Dépenses", "Résultats", "Coût / résultat", "Impressions", "Portée cumulée", "Clics", "CTR (%)", "Type d’action", "Mesure vérifiée", "Insight IDs", "Payload IDs"]
        if combined:
            headers.insert(0, "Compte Meta")
        ws.append(headers)
        for row in snapshot[key]:
            values = [
                row["name"], row["object_id"], _number(row["spend"]), _number(row["results"]), _number(row["cost_per_result"]),
                row["impressions"], row["reach"], row["clicks"], _number(row["ctr"]), row["result_action_type"],
                "Oui" if row["result_verified"] else "Non", ",".join(str(v) for v in row["insight_ids"]), ",".join(str(v) for v in row["raw_payload_ids"]),
            ]
            if combined:
                values.insert(0, row.get("account_name", ""))
            ws.append(values)
        widths = [40, 18, 14, 12, 18, 14, 18, 12, 12, 35, 15, 28, 28]
        if combined:
            widths.insert(0, 30)
        _style_sheet(ws, widths, landscape=True)
        _format_columns(
            ws,
            ({"D": "#,##0.00", "E": "#,##0.00", "F": "#,##0.00", "G": "#,##0", "H": "#,##0", "I": "#,##0", "J": "0.00"} if combined else {"C": "#,##0.00", "D": "#,##0.00", "E": "#,##0.00", "F": "#,##0", "G": "#,##0", "H": "#,##0", "I": "0.00"}),
        )

    alerts = wb.create_sheet("Alertes")
    alert_headers = ["Date", "Gravité", "Règle", "Objet", "Constat", "Description", "Recommandation", "Statut"]
    if combined:
        alert_headers.insert(1, "Compte Meta")
    alerts.append(alert_headers)
    for item in snapshot["anomalies"]:
        values = [item["date"], item["severity"], item["rule_id"], item["name"], item["title"], item["description"], item["recommendation"], item["status"]]
        if combined:
            values.insert(1, item.get("account_name", ""))
        alerts.append(values)
    alert_widths = [14, 12, 28, 36, 34, 58, 58, 18]
    if combined:
        alert_widths.insert(1, 28)
    _style_sheet(alerts, alert_widths, landscape=True)
    red_fill = PatternFill("solid", fgColor="FDE8E4")
    amber_fill = PatternFill("solid", fgColor="FFF4D6")
    severity_column = "C" if combined else "B"
    alerts.conditional_formatting.add(
        f"{severity_column}2:{severity_column}{max(alerts.max_row, 2)}",
        FormulaRule(formula=[f'OR({severity_column}2="critical",{severity_column}2="high")'], fill=red_fill),
    )
    alerts.conditional_formatting.add(
        f"{severity_column}2:{severity_column}{max(alerts.max_row, 2)}",
        FormulaRule(formula=[f'{severity_column}2="medium"'], fill=amber_fill),
    )

    syncs = wb.create_sheet("Synchronisations")
    sync_headers = ["ID", "Créée le", "Début période", "Fin période", "Statut", "Tentative", "Enregistrements", "Pages brutes", "Message"]
    if combined:
        sync_headers.insert(1, "Compte Meta")
    syncs.append(sync_headers)
    for item in snapshot["synchronizations"]:
        values = [item["id"], _localized_timestamp(item["created_at"], snapshot["account"]["timezone"]), item["start"], item["end"], item["status"], item["attempt"], item["records"], item["raw_pages"], item["message"]]
        if combined:
            values.insert(1, item.get("account_name", ""))
        syncs.append(values)
    sync_widths = [10, 21, 14, 14, 15, 12, 16, 14, 56]
    if combined:
        sync_widths.insert(1, 28)
    _style_sheet(syncs, sync_widths, landscape=True)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _trend_drawing(title, labels, values, color, unit):
    width, height = 82 * mm, 46 * mm
    drawing = Drawing(width, height)
    drawing.add(String(7 * mm, height - 5 * mm, title, fontName="Helvetica-Bold", fontSize=8, fillColor=colors.HexColor("#004472")))
    available = [(label, value) for label, value in zip(labels, values) if value is not None]
    if not available:
        drawing.add(String(7 * mm, 19 * mm, "Mesure indisponible", fontName="Helvetica", fontSize=7, fillColor=colors.HexColor("#617487")))
        return drawing
    chart = HorizontalLineChart()
    chart.x = 8 * mm
    chart.y = 8 * mm
    chart.width = 69 * mm
    chart.height = 28 * mm
    chart.data = [[float(value) for _, value in available]]
    chart.categoryAxis.categoryNames = [label for label, _ in available]
    chart.categoryAxis.labels.fontName = "Helvetica"
    chart.categoryAxis.labels.fontSize = 5
    chart.categoryAxis.labels.angle = 30
    chart.categoryAxis.labels.dy = -4
    chart.valueAxis.valueMin = 0
    maximum = max(float(value) for _, value in available)
    chart.valueAxis.valueMax = max(maximum * 1.15, 1)
    chart.valueAxis.labels.fontName = "Helvetica"
    chart.valueAxis.labels.fontSize = 5.5
    chart.lines[0].strokeColor = colors.HexColor(color)
    chart.lines[0].strokeWidth = 1.8
    chart.joinedLines = 1
    drawing.add(chart)
    drawing.add(String(7 * mm, 2 * mm, unit, fontName="Helvetica", fontSize=5.5, fillColor=colors.HexColor("#617487")))
    return drawing


def _page_footer(canvas, doc):
    canvas.saveState()
    page_width, _ = A4
    canvas.setStrokeColor(colors.HexColor("#ffc90d"))
    canvas.setLineWidth(0.7)
    canvas.line(16 * mm, 11 * mm, page_width - 16 * mm, 11 * mm)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#617487"))
    canvas.drawString(16 * mm, 7 * mm, "ULTEx - Rapport Meta interne")
    canvas.drawRightString(page_width - 16 * mm, 7 * mm, f"Page {canvas.getPageNumber()}")
    canvas.restoreState()


def _reportlab_pdf(snapshot: dict) -> bytes:
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=15 * mm, bottomMargin=18 * mm)
    styles = getSampleStyleSheet()
    styles["BodyText"].textColor = colors.HexColor("#0b2d45")
    styles["Title"].fontSize = 22
    styles["Title"].leading = 25
    styles["Title"].spaceAfter = 1
    styles["Title"].textColor = colors.HexColor("#004472")
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=7.7, leading=10, textColor=colors.HexColor("#536b7e")))
    styles.add(ParagraphStyle(name="Footnote", parent=styles["Small"], fontSize=6.5, leading=8, spaceBefore=1))
    styles.add(ParagraphStyle(name="Section", parent=styles["Heading2"], fontSize=13, leading=16, textColor=colors.HexColor("#004472"), spaceBefore=10, spaceAfter=7))
    styles.add(ParagraphStyle(name="Lead", parent=styles["BodyText"], fontSize=10, leading=14, textColor=colors.HexColor("#0b2d45"), borderColor=colors.HexColor("#c9deed"), borderWidth=0.5, borderPadding=8, backColor=colors.HexColor("#f3f8fc")))
    generated_text = _localized_timestamp(snapshot["generated_at"], snapshot["account"]["timezone"])
    combined = bool(snapshot["account"].get("combined"))
    header_text = [
        Paragraph('<font color="#0159a3"><b>META REPORTS</b></font>', styles["Small"]),
        Paragraph("Rapport de performance Meta Ads", styles["Title"]),
        Paragraph(f"{snapshot['period']['start']} au {snapshot['period']['end']} - {snapshot['account']['name']} - {snapshot['account']['currency']}", styles["Small"]),
    ]
    if LOGO_PATH.exists():
        report_logo = RLImage(str(LOGO_PATH), width=45 * mm, height=15 * mm)
        header = Table([[report_logo, header_text]], colWidths=[50 * mm, 120 * mm])
    else:
        header = Table([[header_text]], colWidths=[170 * mm])
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LINEBELOW", (0, 0), (-1, -1), 1.5, colors.HexColor("#ffc90d")),
    ]))
    story = [
        header,
        Spacer(1, 6 * mm),
        Paragraph("1. Cadre du rapport", styles["Section"]),
        Paragraph(f"Généré le {generated_text} - Fuseau {snapshot['account']['timezone']} - Attribution: {snapshot['attribution']['label']}.", styles["BodyText"]),
    ]
    if combined:
        included = ", ".join(f"{item['name']} (act_{item['external_id']})" for item in snapshot["account"].get("accounts", []))
        story.append(Paragraph(f"Comptes inclus: {included}", styles["Small"]))
    story.append(Paragraph("2. État des synchronisations", styles["Section"]))
    if snapshot["data_quality"]:
        for warning in snapshot["data_quality"]:
            story.append(Paragraph(f"- {warning}", styles["Small"]))
    else:
        story.append(Paragraph("Toutes les données requises sont présentes pour la période.", styles["BodyText"]))
    story += [Paragraph("3. Résumé exécutif", styles["Section"]), Paragraph(snapshot["executive_summary"], styles["Lead"]), Paragraph("4. KPI principaux", styles["Section"])]
    k = snapshot["kpis"]
    kpi_table = Table([
        ["Dépenses", "Meta résultats", "Coût / résultat", "Impressions", "CTR"],
        [f"{k['spend']} {snapshot['account']['currency']}" if k["spend"] is not None else "Non consolidé", k["results"] or "Manquant", f"{k['cost_per_result']} {snapshot['account']['currency']}" if k["cost_per_result"] is not None else "Non calculable", str(k["impressions"]), f"{k['ctr'] or 'Non calculable'} %"],
    ], colWidths=[35 * mm] * 5)
    kpi_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#004472")), ("LINEABOVE", (0, 0), (-1, 0), 1.2, colors.HexColor("#ffc90d")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#d8e3ec")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story.append(kpi_table)
    story.append(Paragraph("5. Comparaisons", styles["Section"]))
    comp_data = [["Période", "Dépenses", "Résultats", "Coût / résultat"]]
    for key in ("1", "7", "30"):
        comp = snapshot["comparisons"][key]
        def comparison_value(field):
            value = comp.get(field)
            return f"{value} %" if value is not None else "Indisponible"
        comp_data.append([f"{key} jour(s)", comparison_value("spend_change_percent"), comparison_value("results_change_percent"), comparison_value("cpl_change_percent")])
    comp_table = Table(comp_data, colWidths=[40 * mm] * 4)
    comp_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf3fa")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#004472")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#d8e3ec")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story.append(comp_table)

    def performance_section(number, title, records, ranking=None):
        story.append(Paragraph(f"{number}. {title}", styles["Section"]))
        if ranking:
            strongest = ranking.get("strongest")
            weakest = ranking.get("weakest")
            if strongest:
                story.append(Paragraph(f"Plus efficace par coût/résultat: <b>{strongest['name']}</b> ({strongest['cost_per_result']} {snapshot['account']['currency']}).", styles["Small"]))
            if weakest:
                story.append(Paragraph(f"À examiner par coût/résultat: <b>{weakest['name']}</b> ({weakest['cost_per_result']} {snapshot['account']['currency']}).", styles["Small"]))
        if not records:
            story.append(Paragraph("Aucune donnée disponible.", styles["Small"]))
            return
        data = [["Nom", "Dépenses", "Résultats", "Coût / résultat", "CTR"]]
        for item in records[:12]:
            item_name = item["name"] or item["object_id"]
            if combined:
                item_name = f"{item.get('account_name', '')} - {item_name}"
            data.append([Paragraph(item_name, styles["Small"]), f"{item['spend']} {item.get('currency', snapshot['account']['currency'])}", item["results"] or "-", item["cost_per_result"] or "-", item["ctr"] or "-"])
        table = Table(data, colWidths=[72 * mm, 25 * mm, 22 * mm, 30 * mm, 20 * mm], repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#004472")), ("LINEABOVE", (0, 0), (-1, 0), 1.1, colors.HexColor("#ffc90d")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.4), ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#d8e3ec")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story.append(table)

    performance_section(6, "Performance des campagnes", snapshot["campaigns"], snapshot["rankings"]["campaigns"])
    performance_section(7, "Ensembles et publicités", (snapshot["adsets"][:6] + snapshot["ads"][:6]), snapshot["rankings"]["ads"])
    story.append(Paragraph("8. Tendance", styles["Section"]))
    labels = [item["date"][5:] for item in snapshot["trend"]]
    spend_values = [_number(item["spend"]) for item in snapshot["trend"]]
    result_values = [_number(item["results"]) for item in snapshot["trend"]]
    charts = Table(
        [[
            _trend_drawing("Dépenses", labels, spend_values, "#0159a3", snapshot["account"]["currency"]),
            _trend_drawing("Meta résultats", labels, result_values, "#ffc90d", "résultats"),
        ]],
        colWidths=[84 * mm, 84 * mm],
    )
    charts.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOX", (0, 0), (-1, -1), .35, colors.HexColor("#d8e3ec")), ("INNERGRID", (0, 0), (-1, -1), .35, colors.HexColor("#d8e3ec"))]))
    story.append(charts)
    story.append(Spacer(1, 3 * mm))
    trend_data = [["Date", "Dépenses", "Résultats", "Coût / résultat", "Impressions"]]
    for item in snapshot["trend"][-7:]:
        trend_data.append([item["date"], item["spend"], item["results"] or "-", item["cost_per_result"] or "-", item["impressions"]])
    trend_table = Table(trend_data, colWidths=[33 * mm] * 5, repeatRows=1)
    trend_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf3fa")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#004472")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#d8e3ec")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(trend_table)
    story.append(Paragraph("9. Alertes et vérifications recommandées", styles["Section"]))
    if not snapshot["anomalies"]:
        story.append(Paragraph("Aucune alerte ouverte.", styles["BodyText"]))
    for item in snapshot["anomalies"]:
        story.append(Paragraph(f"[{item['severity'].upper()}] {item['title']} - {item['name']} ({item['rule_id']})", styles["BodyText"]))
        story.append(Paragraph(item["recommendation"], styles["Small"]))
        story.append(Spacer(1, 2 * mm))
    story.append(Paragraph("10. Sources et définitions", styles["Section"]))
    for key, value in snapshot["definitions"].items():
        story.append(Paragraph(f"<b>{key}</b> - {value}", styles["Small"]))
    story.append(Paragraph(f"Source: Meta Marketing API {snapshot['sources']['api_version']} - Insight IDs: {', '.join(str(v) for v in snapshot['sources']['insight_ids'][:50])} - Empreinte payloads: {', '.join(str(v) for v in snapshot['sources']['raw_payload_ids'][:50])}", styles["Footnote"]))
    doc.build(story, onFirstPage=_page_footer, onLaterPages=_page_footer)
    return output.getvalue()


def build_pdf(snapshot: dict) -> bytes:
    html = render_to_string("reporting/report_pdf.html", {"report": snapshot, "logo_data_uri": _logo_data_uri()})
    try:
        from playwright.sync_api import sync_playwright

        with TemporaryDirectory(prefix="ultex-meta-pdf-") as tmp:
            target = Path(tmp) / "report.pdf"
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=True)
                page = browser.new_page(viewport={"width": 1240, "height": 1754})
                page.set_content(html, wait_until="networkidle")
                page.emulate_media(media="print")
                page.pdf(path=str(target), format="A4", print_background=True, margin={"top": "12mm", "right": "12mm", "bottom": "14mm", "left": "12mm"})
                browser.close()
            return target.read_bytes()
    except Exception:
        return _reportlab_pdf(snapshot)
