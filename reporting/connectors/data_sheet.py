from __future__ import annotations

import base64
import json
import re
import time
from dataclasses import dataclass
from datetime import date, datetime, timezone as datetime_timezone
from hashlib import sha256
from pathlib import Path
from threading import Lock
from typing import Any, Iterable
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlencode
from urllib.request import Request, urlopen

from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import padding
from django.conf import settings
from django.core.cache import cache
from django.utils import timezone
from openpyxl import load_workbook


class DataSheetError(RuntimeError):
    """A safe, user-facing Google Sheets connection or parsing error."""


@dataclass(frozen=True)
class Lead:
    source: str
    row_number: int
    entered_on: date
    code: str
    name: str
    contact: str
    channel: str
    product: str
    quantity: str
    status: str
    details: str

    @property
    def source_label(self) -> str:
        return SOURCE_DEFINITIONS.get(self.source, {}).get("label", self.source)

    @property
    def search_text(self) -> str:
        return " ".join(
            (
                self.source_label,
                self.code,
                self.name,
                self.contact,
                self.channel,
                self.product,
                self.quantity,
                self.status,
                self.details,
            )
        ).casefold()


@dataclass(frozen=True)
class LeadDataset:
    leads: tuple[Lead, ...]
    fetched_at: datetime
    mode: str


SOURCE_DEFINITIONS = {
    "landing": {"label": "Landing page", "sheet": "landing page"},
    "clicks": {"label": "Clicks", "sheet": "Clics"},
    "dossiers": {"label": "Dossier IA", "sheet": "Dossiers IA"},
}
CONSOLIDATED_SOURCE = "all"
CONSOLIDATED_LABEL = "Vue consolidée"

LANDING_FIRST_MARKER = date(2025, 6, 23)
CLICKS_FIRST_MARKER = date(2024, 5, 4)
SHEETS_READONLY_SCOPE = "https://www.googleapis.com/auth/spreadsheets.readonly"
_TOKEN_CACHE: dict[str, tuple[str, int]] = {}
_TOKEN_LOCK = Lock()


def source_choices() -> list[tuple[str, str]]:
    return [(CONSOLIDATED_SOURCE, CONSOLIDATED_LABEL)] + [
        (key, value["label"]) for key, value in SOURCE_DEFINITIONS.items()
    ]


def _cell(row: list[Any] | tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Oui" if value else "Non"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).replace("\u00a0", " ").strip()


def _meaningful(value: Any) -> str:
    text = _text(value)
    return "" if text in {"", "-"} else text


def _join_distinct(values: Iterable[str]) -> str:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = _meaningful(value)
        key = cleaned.casefold()
        if cleaned and key not in seen:
            seen.add(key)
            result.append(cleaned)
    return " · ".join(result)


def _parse_day(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    iso_match = re.search(r"\b(\d{4})-(\d{1,2})-(\d{1,2})\b", text)
    if iso_match:
        try:
            return date(*(int(part) for part in iso_match.groups()))
        except ValueError:
            return None
    local_match = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", text)
    if local_match:
        day, month, year = (int(part) for part in local_match.groups())
        try:
            return date(year, month, day)
        except ValueError:
            return None
    return None


def _parse_timestamp(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or datetime_timezone.utc)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time(), tzinfo=datetime_timezone.utc)
    text = _text(value)
    if not text:
        return None
    match = re.search(
        r"\b(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T]+(\d{1,2}):(\d{2})(?::(\d{2}))?)?",
        text,
    )
    if not match:
        parsed_day = _parse_day(text)
        return datetime.combine(parsed_day, datetime.min.time(), tzinfo=datetime_timezone.utc) if parsed_day else None
    year, month, day, hour, minute, second = match.groups()
    try:
        return datetime(
            int(year),
            int(month),
            int(day),
            int(hour or 0),
            int(minute or 0),
            int(second or 0),
            tzinfo=datetime_timezone.utc,
        )
    except ValueError:
        return None


def _latest_label(row: list[Any] | tuple[Any, ...], label_indexes: Iterable[int]) -> str:
    labels = [_meaningful(_cell(row, index)) for index in label_indexes]
    return next((label for label in reversed(labels) if label), "")


def _latest_dossier_status(row: list[Any] | tuple[Any, ...]) -> str:
    candidates: list[tuple[datetime, int, str]] = []
    fallback: list[str] = []
    for date_index in range(14, 34, 2):
        label = _meaningful(_cell(row, date_index + 1))
        if not label:
            continue
        fallback.append(label)
        timestamp = _parse_timestamp(_cell(row, date_index))
        if timestamp:
            candidates.append((timestamp, date_index, label))
    if candidates:
        return max(candidates, key=lambda item: (item[0], item[1]))[2]
    return fallback[-1] if fallback else ""


def parse_landing_rows(rows: Iterable[list[Any] | tuple[Any, ...]]) -> list[Lead]:
    leads: list[Lead] = []
    current_day: date | None = None
    for row_number, row in enumerate(rows, start=1):
        code = _meaningful(_cell(row, 1))
        if not code or code.casefold() == "code client":
            marker = _parse_day(_cell(row, 0))
            marker_text = _text(_cell(row, 0)).casefold()
            if marker and ("le" in marker_text or isinstance(_cell(row, 0), (date, datetime))):
                current_day = marker if marker >= LANDING_FIRST_MARKER else None
            continue
        if current_day is None:
            continue
        leads.append(
            Lead(
                source="landing",
                row_number=row_number,
                entered_on=current_day,
                code=code,
                name=_meaningful(_cell(row, 3)),
                contact=_meaningful(_cell(row, 4)),
                channel=_meaningful(_cell(row, 2)),
                product=_meaningful(_cell(row, 6)),
                quantity=_meaningful(_cell(row, 9)),
                status=_latest_label(row, (13, 15, 17, 19, 21)),
                details=_join_distinct((_cell(row, 5), _cell(row, 11))),
            )
        )
    return leads


def parse_click_rows(rows: Iterable[list[Any] | tuple[Any, ...]]) -> list[Lead]:
    leads: list[Lead] = []
    current_day: date | None = None
    for row_number, row in enumerate(rows, start=1):
        code = _meaningful(_cell(row, 1))
        if not code or code.casefold() == "code client":
            marker = _parse_day(_cell(row, 0))
            if marker:
                current_day = marker if marker >= CLICKS_FIRST_MARKER else None
            continue
        if current_day is None:
            continue
        form_value = _meaningful(_cell(row, 2))
        observation = _meaningful(_cell(row, 9))
        leads.append(
            Lead(
                source="clicks",
                row_number=row_number,
                entered_on=current_day,
                code=code,
                name="",
                contact=_meaningful(_cell(row, 4)),
                channel=_meaningful(_cell(row, 3)),
                product=_meaningful(_cell(row, 5)),
                quantity=_meaningful(_cell(row, 6)),
                status=_latest_label(row, (11, 13, 15, 17, 19)),
                details=_join_distinct((f"Formulaire {form_value}" if form_value else "", observation)),
            )
        )
    return leads


def parse_dossier_rows(rows: Iterable[list[Any] | tuple[Any, ...]]) -> list[Lead]:
    leads: list[Lead] = []
    for row_number, row in enumerate(rows, start=1):
        code = _meaningful(_cell(row, 1))
        if not code or code.casefold() == "code client":
            continue
        entered_on = _parse_day(_cell(row, 14))
        if entered_on is None:
            continue
        country = _meaningful(_cell(row, 4))
        origin = _meaningful(_cell(row, 9))
        leads.append(
            Lead(
                source="dossiers",
                row_number=row_number,
                entered_on=entered_on,
                code=code,
                name=_meaningful(_cell(row, 2)),
                contact=_meaningful(_cell(row, 3)),
                channel=_meaningful(_cell(row, 5)),
                product=_meaningful(_cell(row, 6)),
                quantity=_meaningful(_cell(row, 7)),
                status=_latest_dossier_status(row),
                details=_join_distinct((country, origin)),
            )
        )
    return leads


PARSERS = {
    "landing": parse_landing_rows,
    "clicks": parse_click_rows,
    "dossiers": parse_dossier_rows,
}


def _base64url(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).rstrip(b"=").decode("ascii")


def _service_account_info() -> dict[str, Any]:
    file_path = str(getattr(settings, "GOOGLE_SERVICE_ACCOUNT_FILE", "") or "").strip()
    inline_json = str(getattr(settings, "GOOGLE_SERVICE_ACCOUNT_JSON", "") or "").strip()
    encoded_json = str(getattr(settings, "GOOGLE_SERVICE_ACCOUNT_JSON_BASE64", "") or "").strip()
    try:
        if file_path:
            return json.loads(Path(file_path).read_text(encoding="utf-8"))
        if inline_json:
            return json.loads(inline_json)
        if encoded_json:
            return json.loads(base64.b64decode(encoded_json).decode("utf-8"))
    except (OSError, ValueError, UnicodeDecodeError) as exc:
        raise DataSheetError("Les identifiants Google Sheets sont illisibles ou incomplets.") from exc
    raise DataSheetError(
        "Connexion Google Sheets non configurée. Ajoutez un compte de service en lecture seule pour afficher les leads en direct."
    )


def _access_token(service_account: dict[str, Any]) -> str:
    client_email = _meaningful(service_account.get("client_email"))
    private_key = _meaningful(service_account.get("private_key"))
    token_uri = _meaningful(service_account.get("token_uri")) or "https://oauth2.googleapis.com/token"
    if not client_email or not private_key:
        raise DataSheetError("Le compte de service Google Sheets ne contient pas les informations requises.")
    now = int(time.time())
    with _TOKEN_LOCK:
        cached = _TOKEN_CACHE.get(client_email)
        if cached and cached[1] > now + 60:
            return cached[0]

        header = _base64url(json.dumps({"alg": "RS256", "typ": "JWT"}, separators=(",", ":")).encode("utf-8"))
        claims = _base64url(
            json.dumps(
                {
                    "iss": client_email,
                    "scope": SHEETS_READONLY_SCOPE,
                    "aud": token_uri,
                    "iat": now,
                    "exp": now + 3600,
                },
                separators=(",", ":"),
            ).encode("utf-8")
        )
        signing_input = f"{header}.{claims}".encode("ascii")
        try:
            key = serialization.load_pem_private_key(private_key.encode("utf-8"), password=None)
            signature = key.sign(signing_input, padding.PKCS1v15(), hashes.SHA256())
        except (TypeError, ValueError) as exc:
            raise DataSheetError("La clé privée du compte de service Google Sheets est invalide.") from exc
        assertion = f"{header}.{claims}.{_base64url(signature)}"
        request = Request(
            token_uri,
            data=urlencode({"grant_type": "urn:ietf:params:oauth:grant-type:jwt-bearer", "assertion": assertion}).encode("ascii"),
            headers={"Content-Type": "application/x-www-form-urlencoded"},
            method="POST",
        )
        timeout = int(getattr(settings, "GOOGLE_SHEETS_TIMEOUT_SECONDS", 15))
        try:
            with urlopen(request, timeout=timeout) as response:
                payload = json.loads(response.read().decode("utf-8"))
        except (HTTPError, URLError, TimeoutError, ValueError) as exc:
            raise DataSheetError("Google n’a pas pu authentifier la connexion en lecture seule.") from exc
        token = _meaningful(payload.get("access_token"))
        if not token:
            raise DataSheetError("Google n’a retourné aucun jeton d’accès pour la feuille.")
        expires_in = int(payload.get("expires_in") or 3600)
        _TOKEN_CACHE[client_email] = (token, now + expires_in)
        return token


def _fetch_google_rows(sheet_name: str) -> list[list[Any]]:
    spreadsheet_id = _meaningful(getattr(settings, "GOOGLE_SHEET_ID", ""))
    if not spreadsheet_id:
        raise DataSheetError("L’identifiant du classeur Google Sheets n’est pas configuré.")
    token = _access_token(_service_account_info())
    range_name = quote(f"'{sheet_name}'!A:BC", safe="")
    query = urlencode(
        {
            "majorDimension": "ROWS",
            "valueRenderOption": "FORMATTED_VALUE",
            "dateTimeRenderOption": "FORMATTED_STRING",
        }
    )
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{quote(spreadsheet_id, safe='')}/values/{range_name}?{query}"
    request = Request(url, headers={"Authorization": f"Bearer {token}", "Accept": "application/json"})
    timeout = int(getattr(settings, "GOOGLE_SHEETS_TIMEOUT_SECONDS", 15))
    try:
        with urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        if exc.code == 403:
            message = "Le compte de service n’a pas accès à ce classeur. Partagez-le avec son adresse e-mail en mode Lecteur."
        elif exc.code == 404:
            message = "Le classeur Google Sheets ou la feuille demandée est introuvable."
        else:
            message = "Google Sheets est momentanément indisponible. Réessayez dans un instant."
        raise DataSheetError(message) from exc
    except (URLError, TimeoutError, ValueError) as exc:
        raise DataSheetError("Impossible de joindre Google Sheets pour le moment.") from exc
    values = payload.get("values")
    if not isinstance(values, list):
        raise DataSheetError("Google Sheets a retourné une réponse inattendue.")
    return [list(row) for row in values if isinstance(row, list)]


def _load_local_rows(workbook_path: str, sheet_name: str) -> list[list[Any]]:
    path = Path(workbook_path)
    if not path.is_file():
        raise DataSheetError("Le fichier DATA local configuré est introuvable.")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise DataSheetError(f"La feuille « {sheet_name} » est absente du fichier DATA local.")
            return [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
        finally:
            workbook.close()
    except DataSheetError:
        raise
    except (OSError, ValueError) as exc:
        raise DataSheetError("Le fichier DATA local ne peut pas être lu.") from exc


def load_lead_dataset(source: str, *, refresh: bool = False) -> LeadDataset:
    if source == CONSOLIDATED_SOURCE:
        datasets: list[LeadDataset] = []
        for source_key, definition in SOURCE_DEFINITIONS.items():
            try:
                datasets.append(load_lead_dataset(source_key, refresh=refresh))
            except DataSheetError as exc:
                raise DataSheetError(f"{definition['label']} : {exc}") from exc
        return LeadDataset(
            leads=tuple(lead for dataset in datasets for lead in dataset.leads),
            fetched_at=max(dataset.fetched_at for dataset in datasets),
            mode=datasets[0].mode,
        )
    if source not in SOURCE_DEFINITIONS:
        raise DataSheetError("La feuille demandée n’existe pas.")
    definition = SOURCE_DEFINITIONS[source]
    workbook_path = str(getattr(settings, "DATA_SHEET_LOCAL_XLSX", "") or "").strip()
    spreadsheet_id = _meaningful(getattr(settings, "GOOGLE_SHEET_ID", ""))
    cache_identity = workbook_path or spreadsheet_id
    if workbook_path:
        try:
            cache_identity = f"{workbook_path}:{Path(workbook_path).stat().st_mtime_ns}"
        except OSError:
            pass
    cache_hash = sha256(cache_identity.encode("utf-8")).hexdigest()[:16]
    cache_key = f"data-sheet:v1:{cache_hash}:{source}"
    if not refresh:
        cached = cache.get(cache_key)
        if isinstance(cached, LeadDataset):
            return cached

    rows = _load_local_rows(workbook_path, definition["sheet"]) if workbook_path else _fetch_google_rows(definition["sheet"])
    dataset = LeadDataset(
        leads=tuple(PARSERS[source](rows)),
        fetched_at=timezone.now(),
        mode="local" if workbook_path else "google",
    )
    cache_seconds = max(int(getattr(settings, "GOOGLE_SHEETS_CACHE_SECONDS", 60)), 0)
    cache.set(cache_key, dataset, timeout=cache_seconds)
    return dataset
